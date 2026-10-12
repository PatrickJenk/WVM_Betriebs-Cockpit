import sys, re
from pathlib import Path
sys.path.insert(0,str(Path('tmp/manual-deps').resolve()))
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

out=Path('output/handbuch');out.mkdir(parents=True,exist_ok=True)
doc=Document();sec=doc.sections[0]
sec.page_width=Cm(21);sec.page_height=Cm(29.7)
sec.top_margin=Cm(1.4);sec.bottom_margin=Cm(1.4)
sec.left_margin=Cm(1.8);sec.right_margin=Cm(1.8)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3']:
 st=doc.styles[name];st.font.name='Calibri';st.font.color.rgb=RGBColor(0,0,0)
 st.paragraph_format.space_after=Pt(4)
doc.styles['Normal'].font.size=Pt(10)
doc.styles['Normal'].paragraph_format.line_spacing=1.0
doc.styles['Title'].font.size=Pt(22)
doc.styles['Heading 1'].font.size=Pt(19)
doc.styles['Heading 2'].font.size=Pt(12)
doc.styles['Heading 2'].paragraph_format.space_before=Pt(7)
header=sec.header.paragraphs[0];header.text='Wärmeverbund Betriebs Cockpit  |  Benutzerhandbuch'
header.style=doc.styles['Normal'];header.runs[0].font.size=Pt(8)
footer=sec.footer.paragraphs[0];footer.text='Stand 07.10.2026  ·  Version 1.0                                      Seite '
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)
for r in footer.runs:r.font.size=Pt(8)

def paragraph(text,style=None):
 p=doc.add_paragraph(style=style)
 # Keep inline backticks and bold as readable prose in this manual.
 p.add_run(text.replace('**','').replace('`',''))
 return p

lines=Path('HANDBUCH.md').read_text(encoding='utf-8').splitlines();i=0
while i<len(lines):
 line=lines[i].strip()
 if not line:i+=1;continue
 if line=='<!-- PAGE -->':doc.add_page_break();i+=1;continue
 if line.startswith('|'):
  rows=[]
  while i<len(lines) and lines[i].strip().startswith('|'):
   cells=[c.strip() for c in lines[i].strip().strip('|').split('|')]
   if not all(re.fullmatch(r'[- :]+',c) for c in cells):rows.append(cells)
   i+=1
  table=doc.add_table(rows=0,cols=len(rows[0]));table.autofit=False
  widths=([8.9,8.5] if len(rows[0])==2 else [9.3,4.2,3.9])
  if rows[0][0]=='Seite':widths=[2,15.4]
  if rows[0][0]=='Anzeige':widths=[4.6,3.3,9.5]
  for col,w in zip(table.columns,widths):col.width=Cm(w)
  borders=OxmlElement('w:tblBorders')
  for edge in ['top','left','bottom','right','insideH','insideV']:
   el=OxmlElement('w:'+edge);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');borders.append(el)
  table._tbl.tblPr.append(borders)
  for j,values in enumerate(rows):
   row=table.add_row();pr=row._tr.get_or_add_trPr();pr.append(OxmlElement('w:cantSplit'))
   if j==0:pr.append(OxmlElement('w:tblHeader'))
   for k,(cell,value) in enumerate(zip(row.cells,values)):
    cell.width=Cm(widths[k]);cell.vertical_alignment=1
    tcpr=cell._tc.get_or_add_tcPr();shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'12384C' if j==0 else ('F2F6F8' if j%2 else 'FFFFFF'));tcpr.append(shade)
    margin=OxmlElement('w:tcMar')
    for side in ['top','bottom','left','right']:
     el=OxmlElement('w:'+side);el.set(qn('w:w'),'65');el.set(qn('w:type'),'dxa');margin.append(el)
    tcpr.append(margin);p=cell.paragraphs[0];p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1
    run=p.add_run(value);run.font.size=Pt(9.5)
    if j==0:run.bold=True;run.font.color.rgb=RGBColor(255,255,255)
    if k==1 and len(values)==3:
     color={'Grün':'37835A','Blau':'0078BF','Rot':'C62828','Gelb':'987900','Gelb bis Orange':'B85B00','Orange':'B85B00'}.get(value)
     if color and j:run.font.color.rgb=RGBColor.from_string(color);run.bold=True
  doc.add_paragraph().paragraph_format.space_after=Pt(0)
  continue
 if line.startswith('# '):paragraph(line[2:],'Title')
 elif line.startswith('## '):paragraph(line[3:],'Heading 1')
 elif line.startswith('### '):paragraph(line[4:],'Heading 2')
 else:
  # Join ordinary wrapped prose, keeping numbered instructions separate.
  if not re.match(r'^\d+\.',line):
   while i+1<len(lines) and lines[i+1].strip() and not lines[i+1].strip().startswith(('#','|','<!--')) and not re.match(r'^\d+\.',lines[i+1].strip()):
    i+=1;line+=' '+lines[i].strip()
  p=paragraph(line)
  if re.match(r'^\d+\.',line):p.paragraph_format.space_after=Pt(4)
 i+=1
doc.core_properties.title='Benutzerhandbuch Wärmeverbund Betriebs Cockpit'
doc.core_properties.subject='Bedienung, Kommunikation und Benutzerrechte'
for st in doc.styles:
 for border in list(st.element.iter(qn('w:pBdr'))): border.getparent().remove(border)
doc.save(out/'Benutzerhandbuch_Waermeverbund_Cockpit.docx')
print('DOCX erstellt')
